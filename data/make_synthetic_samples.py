"""
Generates two SYNTHETIC sample project plans (Project C and Project D) using
the same column schema as the real Project_Plan_B.xlsx export, so that
Phase 3 (monthly synthesis across a portfolio) has more than one project to
compare. These are clearly fictional and used only to demonstrate the
portfolio/trend-analysis capability -- see README for how to swap in real
files.

Project C: consistently healthy (Green) project, on schedule.
Project D: a deteriorating project (Amber sliding to Red) with PM status
           comments included, to demonstrate the sentiment-proxy dimension.
"""
import openpyxl
from datetime import date, timedelta
import random

HEADERS = ['No.of days Until Today', 'No.of days', 'Target start date to Today',
           'Project Name', 'Project Category', 'Ancestors', 'Project Manager',
           'Phase/Milestone', 'Area', 'At Risk?', 'Schedule Health', 'Task Name',
           'Status', '% Complete', 'Start Date', 'End Date', 'Priority', 'Owner',
           'On Hold?', 'Not Applicable?', 'Duration', 'Predecessors', 'Total Float',
           'Critical ?', 'Baseline Start', 'Baseline Finish', 'Variance',
           'Status Comment', 'Baseline Start Date', 'Baseline End Date', 'Comments',
           'Assigned To', 'Start', 'Finish', 'Baseline Start2', 'Baseline Finish2',
           'Variance2']


def blank_row():
    return [None] * len(HEADERS)


def add_task(rows, name, status, pct, start, end, ancestors, health="Green",
             critical=None, predecessors=None, on_hold=None, comment=None):
    r = blank_row()
    r[HEADERS.index('Ancestors')] = ancestors
    r[HEADERS.index('Task Name')] = name
    r[HEADERS.index('Status')] = status
    r[HEADERS.index('% Complete')] = pct
    r[HEADERS.index('Start Date')] = start
    r[HEADERS.index('End Date')] = end
    r[HEADERS.index('Schedule Health')] = health
    r[HEADERS.index('Critical ?')] = critical
    r[HEADERS.index('Predecessors')] = predecessors
    r[HEADERS.index('On Hold?')] = on_hold
    r[HEADERS.index('Status Comment')] = comment
    rows.append(r)


def build(path, project_name, pm, rows_spec, today, at_risk, sched_health, stage):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Project Plan"
    ws.append(HEADERS)
    rows = []
    # root row (ancestors=0) carries the project name, matching real export convention
    add_task(rows, project_name, "In Progress", None, None, None, 0)
    for spec in rows_spec:
        add_task(rows, *spec)
    for r in rows:
        ws.append(r)

    wb.create_sheet("Comments")

    sm = wb.create_sheet("Summary")
    pct_vals = [s[2] for s in rows_spec if isinstance(s[2], (int, float))]
    overall_pct = sum(pct_vals) / len(pct_vals) if pct_vals else 0
    starts = [s[3] for s in rows_spec if s[3]]
    ends = [s[4] for s in rows_spec if s[4]]
    summary_kv = [
        ("Project Name", project_name),
        ("Project Manager", pm),
        ("Project Start Date", min(starts) if starts else None),
        ("Project End Date", max(ends) if ends else None),
        ("Not Started", sum(1 for s in rows_spec if s[1] == "Not Started")),
        ("In Progress", sum(1 for s in rows_spec if s[1] == "In Progress")),
        ("Completed", sum(1 for s in rows_spec if s[1] == "Completed")),
        ("On Hold", sum(1 for s in rows_spec if s[1] == "On Hold")),
        ("At Risk", at_risk),
        ("Project Stage", stage),
        ("% Complete", round(overall_pct, 2)),
        ("Schedule Health", sched_health),
        ("Today's Date", today),
        ("Project Status", "In Progress"),
    ]
    for k, v in summary_kv:
        sm.append([k, v])

    wb.save(path)
    print("wrote", path)


TODAY = date(2026, 7, 2)


def d(days_from_today):
    return TODAY + timedelta(days=days_from_today)


# ---------------------------------------------------------------------------
# Project C -- healthy, on-track SaaS rollout
# ---------------------------------------------------------------------------
project_c_tasks = [
    ("Contract Sign Off", "Completed", 1.0, d(-120), d(-118), 1, "Green", True, None, None, None),
    ("Discovery & Requirements", "Completed", 1.0, d(-115), d(-95), 1, "Green", True, None, None, None),
    ("Environment Setup", "Completed", 1.0, d(-94), d(-85), 1, "Green", None, "2", None, None),
    ("Core Configuration", "Completed", 1.0, d(-84), d(-40), 1, "Green", True, "3", None, None),
    ("Integration Build", "In Progress", 0.8, d(-39), d(5), 1, "Green", True, "4", None,
     "Tracking to plan, team is confident in the current timeline."),
    ("UAT Planning", "In Progress", 0.5, d(-10), d(10), 1, "Green", None, "5", None, None),
    ("UAT Execution", "Not Started", 0.0, d(11), d(30), 1, "Green", None, "6", None, None),
    ("Go-Live Readiness", "Not Started", 0.0, d(31), d(45), 1, "Green", True, "7", None, None),
    ("Hypercare", "Not Started", 0.0, d(46), d(60), 1, "Green", None, "8", None, None),
]

# ---------------------------------------------------------------------------
# Project D -- deteriorating implementation, slipping badly
# ---------------------------------------------------------------------------
project_d_tasks = [
    ("Contract Sign Off", "Completed", 1.0, d(-150), d(-148), 1, "Green", True, None, None, None),
    ("Discovery & Requirements", "Completed", 1.0, d(-147), d(-120), 1, "Yellow", True, None, None,
     "Scope grew mid-phase, client kept adding requirements."),
    ("Data Migration Design", "Completed", 1.0, d(-119), d(-95), 1, "Yellow", True, "2", None,
     "Design finalized late; client SME availability was a recurring issue."),
    ("Data Migration Build", "In Progress", 0.35, d(-94), d(-30), 1, "Red", True, "3", None,
     "Significantly behind, two key resources reassigned off the project."),
    ("Integration Build", "In Progress", 0.2, d(-60), d(-10), 1, "Red", True, "3", None,
     "Blocked waiting on client API credentials for three weeks."),
    ("Security Review", "Not Started", 0.0, d(-20), d(-5), 1, "Red", True, "5", None,
     "Cannot start until integration build stabilizes; client escalating."),
    ("UAT Planning", "Not Started", 0.0, d(-5), d(15), 1, "Red", None, "4", None,
     "On hold pending scope re-baseline conversation with sponsor."),
    ("UAT Execution", "Not Started", 0.0, d(16), d(35), 1, "Yellow", None, "7", None, None),
    ("Go-Live Readiness", "Not Started", 0.0, d(36), d(50), 1, "Yellow", True, "8", None, None),
]

build("Project_Plan_C_Sample.xlsx", "Meridian HR Suite Rollout - Northwind Retail",
      "Ayesha Kulkarni", project_c_tasks, TODAY, at_risk="Low", sched_health="Green",
      stage="Build Phase")

build("Project_Plan_D_Sample.xlsx", "Orion ERP Migration - Falcon Manufacturing",
      "Marcus Webb", project_d_tasks, TODAY, at_risk="High", sched_health="Red",
      stage="Build Phase")
