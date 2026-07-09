"""
Robust reader for messy project-plan spreadsheet exports.

Design goals (this is the "handles incomplete/messy data gracefully"
requirement from the assignment):
  - Never crash on a missing column -- just mark that signal unavailable.
  - Never crash on a bad cell value (e.g. "#UNPARSEABLE", blank, wrong type).
  - Always return a structured, typed dataset plus a list of data-quality
    warnings so the agent can be honest about what it couldn't evaluate.
"""

from dataclasses import dataclass, field
from datetime import datetime, date
from typing import Optional, Any
import re
import openpyxl

from .config import COLUMN_ALIASES


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", str(s).strip().lower())


def _to_bool(v: Any) -> Optional[bool]:
    if isinstance(v, bool):
        return v
    if v is None:
        return None
    s = _norm(v)
    if s in ("true", "yes", "y", "1"):
        return True
    if s in ("false", "no", "n", "0", ""):
        return False
    return None


def _to_float(v: Any) -> Optional[float]:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s or s.upper() in ("#UNPARSEABLE", "N/A", "NA", "#N/A", "#VALUE!"):
        return None
    s = s.replace("%", "").replace(",", "")
    try:
        val = float(s)
        return val / 100 if "%" in str(v) else val
    except ValueError:
        return None


def _to_date(v: Any) -> Optional[date]:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    if not s or s.upper() in ("#UNPARSEABLE", "N/A", "NA"):
        return None
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%d-%b-%Y", "%b %d, %Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _to_str(v: Any) -> Optional[str]:
    if v is None:
        return None
    s = str(v).strip()
    if not s or s.upper() in ("#UNPARSEABLE", "N/A", "NA"):
        return None
    return s


def _to_int(v: Any) -> Optional[int]:
    f = _to_float(v)
    return int(f) if f is not None else None


@dataclass
class Task:
    task_name: Optional[str] = None
    status: Optional[str] = None
    pct_complete: Optional[float] = None
    start_date: Optional[date] = None
    end_date: Optional[date] = None
    baseline_start: Optional[date] = None
    baseline_end: Optional[date] = None
    schedule_health: Optional[str] = None
    priority: Optional[str] = None
    critical: Optional[bool] = None
    total_float: Optional[float] = None
    on_hold: Optional[bool] = None
    not_applicable: Optional[bool] = None
    predecessors: Optional[str] = None
    ancestors: Optional[int] = None  # outline depth; 1 = top-level phase
    owner: Optional[str] = None
    status_comment: Optional[str] = None
    area: Optional[str] = None
    phase_milestone: Optional[str] = None
    raw: dict = field(default_factory=dict)


@dataclass
class ProjectPlan:
    project_name: str
    project_manager: Optional[str]
    tasks: list
    warnings: list
    available_fields: set
    source_file: str
    as_of: date


def _build_header_map(header_row) -> dict:
    """Map canonical field name -> column index, based on aliases."""
    norm_headers = {_norm(h): i for i, h in enumerate(header_row) if h is not None}
    field_to_idx = {}
    for canonical, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            if alias in norm_headers:
                field_to_idx[canonical] = norm_headers[alias]
                break
    return field_to_idx


def load_project_plan(path: str, as_of: Optional[date] = None) -> ProjectPlan:
    warnings = []
    as_of = as_of or date.today()

    wb = openpyxl.load_workbook(path, data_only=True)
    sheet_name = "Project Plan" if "Project Plan" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"{path}: sheet '{sheet_name}' is empty")

    header_map = _build_header_map(rows[0])
    required_missing = [f for f in ("status", "task_name") if f not in header_map]
    if required_missing:
        warnings.append(
            f"Missing critical column(s) {required_missing} in '{sheet_name}'; "
            "task-level status inference will be unreliable."
        )

    def get(row, field_name):
        idx = header_map.get(field_name)
        return row[idx] if idx is not None and idx < len(row) else None

    tasks = []
    for row in rows[1:]:
        if row is None or all(c is None for c in row):
            continue
        t = Task(
            task_name=_to_str(get(row, "task_name")),
            status=_to_str(get(row, "status")),
            pct_complete=_to_float(get(row, "pct_complete")),
            start_date=_to_date(get(row, "start_date")),
            end_date=_to_date(get(row, "end_date")),
            baseline_start=_to_date(get(row, "baseline_start")),
            baseline_end=_to_date(get(row, "baseline_end")),
            schedule_health=_to_str(get(row, "schedule_health")),
            priority=_to_str(get(row, "priority")),
            critical=_to_bool(get(row, "critical")),
            total_float=_to_float(get(row, "total_float")),
            on_hold=_to_bool(get(row, "on_hold")),
            not_applicable=_to_bool(get(row, "not_applicable")),
            predecessors=_to_str(get(row, "predecessors")),
            ancestors=_to_int(get(row, "ancestors")),
            owner=_to_str(get(row, "owner")),
            status_comment=_to_str(get(row, "status_comment")),
            area=_to_str(get(row, "area")),
            phase_milestone=_to_str(get(row, "phase_milestone")),
        )
        if t.task_name is None and t.status is None:
            continue  # not a usable row
        tasks.append(t)

    if not tasks:
        warnings.append("No usable task rows found after cleaning.")

    # Project name / PM: prefer a "Summary" sheet if present, else first task row.
    project_name = None
    project_manager = None
    if "Summary" in wb.sheetnames:
        sm = wb["Summary"]
        kv = {}
        for r in sm.iter_rows(values_only=True):
            if r and r[0]:
                kv[_norm(r[0])] = r[1] if len(r) > 1 else None
        project_name = _to_str(kv.get("project name"))
        project_manager = _to_str(kv.get("project manager"))

    if not project_name:
        for row in rows[1:]:
            v = _to_str(get(row, "project_name"))
            if v:
                project_name = v
                break
    if not project_name:
        # Some PM tool exports (e.g. Smartsheet/MS Project) put the project's
        # own name in the root outline row (Ancestors/indent level 0) rather
        # than a dedicated Project Name column.
        for row in rows[1:]:
            if _to_int(get(row, "ancestors")) == 0:
                v = _to_str(get(row, "task_name"))
                if v:
                    project_name = v
                    break
    if not project_manager:
        for row in rows[1:]:
            v = _to_str(get(row, "project_manager"))
            if v:
                project_manager = v
                break

    if not project_name:
        import os
        project_name = os.path.splitext(os.path.basename(path))[0]
        warnings.append(
            f"Project name not found in data; falling back to filename '{project_name}'."
        )

    available_fields = set(header_map.keys())
    for f in ("area", "phase_milestone", "status_comment", "priority", "on_hold"):
        if f not in available_fields:
            warnings.append(f"Field '{f}' not present in source file -- related signal skipped.")

    return ProjectPlan(
        project_name=project_name,
        project_manager=project_manager,
        tasks=tasks,
        warnings=warnings,
        available_fields=available_fields,
        source_file=path,
        as_of=as_of,
    )
